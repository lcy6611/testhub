"""Skill 多格式导出：把生成结果（Markdown 表格）转换为 Excel / 飞书思维导图 / XMind。

核心能力（对齐公众号 universal-testcase-generator）：
- 解析模型输出的 Markdown 表格（表头 + 数据行）
- 按团队模板列（template_columns）做智能字段映射（别名匹配），缺失列留空
- 输出三种格式：Excel（按模板列）、飞书思维导图（Markdown）、XMind（.xmind）
"""

from __future__ import annotations

import io
import json
import re
import zipfile
import uuid

# 列名别名映射：模型常见叫法 -> 模板标准列名（用于智能字段映射）
_COLUMN_ALIASES = {
    "用例编号": "用例编号", "用例id": "用例编号", "用例id ": "用例编号", "id": "用例编号", "编号": "用例编号",
    "用例标题": "用例标题", "标题": "用例标题", "测试点": "用例标题", "用例名称": "用例标题",
    "模块": "模块", "所属模块": "模块", "功能模块": "模块", "模块名称": "模块",
    "优先级": "优先级", "优先级别": "优先级", "等级": "优先级",
    "用例类型": "用例类型", "类型": "用例类型",
    "前置条件": "前置条件", "前置操作": "前置条件", "前置": "前置条件",
    "测试步骤": "测试步骤", "步骤": "测试步骤", "操作步骤": "测试步骤", "操作": "测试步骤",
    "测试数据": "测试数据", "数据": "测试数据",
    "预期结果": "预期结果", "期望结果": "预期结果", "预期": "预期结果",
    "设计者": "设计者", "作者": "设计者", "编写人": "设计者",
    "用例描述": "用例描述", "描述": "用例描述",
}


def _norm(col: str) -> str:
    return (col or "").strip().lower()


def _alias_of(col: str) -> str:
    """返回列名对应的标准名（别名->标准），没有别名则返回原名。"""
    n = _norm(col)
    return _COLUMN_ALIASES.get(n, col.strip())


def parse_markdown_table(md_text: str):
    """从 Markdown 文本中提取第一个表格，返回 (headers, rows)。

    headers: List[str]
    rows:   List[List[str]]
    未找到表格时返回 ([], [])。
    """
    if not md_text:
        return [], []
    lines = md_text.splitlines()
    table_lines = []
    in_table = False
    for line in lines:
        if line.strip().startswith("|"):
            in_table = True
            table_lines.append(line.strip())
        elif in_table:
            # 表格结束（遇到非 | 开头行）
            break
    if len(table_lines) < 2:
        return [], []

    def _split(row_line: str):
        # 去掉首尾 |
        s = row_line.strip()
        if s.startswith("|"):
            s = s[1:]
        if s.endswith("|"):
            s = s[:-1]
        return [c.strip() for c in s.split("|")]

    headers = _split(table_lines[0])
    # 跳过分隔行（第二行如 |---|---|）
    data_start = 1
    if len(table_lines) > 1 and re.match(r"^[\s:|-]+$", table_lines[1].replace("|", "").strip()):
        data_start = 2
    rows = []
    for rl in table_lines[data_start:]:
        cells = _split(rl)
        if any(c for c in cells):
            # 补齐到与表头相同长度
            if len(cells) < len(headers):
                cells = cells + [""] * (len(headers) - len(cells))
            rows.append(cells)
    return headers, rows


def map_to_template(headers, rows, template_columns):
    """把模型输出的 (headers, rows) 映射到模板列顺序。

    返回 (out_headers, out_rows)。
    - template_columns 为空：原样返回（保持模型输出列顺序）
    - 否则：out_headers = template_columns；按别名匹配模型列，缺失列填空
    """
    if not template_columns:
        return headers, rows

    # 建立 模型列标准名 -> 索引
    std_to_idx = {}
    for i, h in enumerate(headers):
        std = _alias_of(h)
        if std not in std_to_idx:
            std_to_idx[std] = i

    out_headers = list(template_columns)
    out_rows = []
    for row in rows:
        out_row = []
        for tcol in out_headers:
            std = _alias_of(tcol)
            idx = std_to_idx.get(std)
            if idx is not None and idx < len(row):
                out_row.append(row[idx])
            else:
                out_row.append("")
        out_rows.append(out_row)
    return out_headers, out_rows


def build_excel_bytes(headers, rows) -> bytes:
    """生成 xlsx 文件字节（openpyxl）。"""
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "测试用例"
    if headers:
        ws.append(headers)
        # 表头样式
        for col_idx in range(1, len(headers) + 1):
            c = ws.cell(row=1, column=col_idx)
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="4472C4")
            c.alignment = Alignment(horizontal="center", vertical="center")
    for r in rows:
        ws.append(r)
    # 简单列宽
    for col_idx in range(1, len(headers) + 1 if headers else 1):
        ws.column_dimensions[ws.cell(row=1, column=col_idx).column_letter if headers else "A"].width = 18
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_feishu_markdown(headers, rows, title: str) -> str:
    """生成飞书思维导图友好的 Markdown（嵌套列表）。"""
    lines = [f"# {title or '测试用例'}", ""]
    for i, r in enumerate(rows, 1):
        # 用例主题：优先用例标题/编号，否则 用例N
        subject = ""
        for key in ("用例标题", "标题", "测试点", "用例编号", "编号"):
            if key in headers:
                v = r[headers.index(key)]
                if v:
                    subject = v
                    break
        if not subject:
            subject = f"用例{i}"
        lines.append(f"- {subject}")
        for h, v in zip(headers, r):
            if v and h not in ("用例标题", "标题", "测试点"):
                lines.append(f"  - {h}：{v}")
        lines.append("")
    return "\n".join(lines)


def build_xmind_bytes(headers, rows, title: str) -> bytes:
    """生成 .xmind 文件字节（zip + content.json）。"""
    def _topic(t: str, children=None):
        node = {
            "id": uuid.uuid4().hex[:12],
            "title": t,
        }
        if children:
            node["children"] = {"attached": children}
        return node

    root_children = []
    for i, r in enumerate(rows, 1):
        subject = ""
        for key in ("用例标题", "标题", "测试点", "用例编号", "编号"):
            if key in headers:
                v = r[headers.index(key)]
                if v:
                    subject = v
                    break
        if not subject:
            subject = f"用例{i}"
        sub = [_topic(f"{h}：{v}") for h, v in zip(headers, r) if v]
        root_children.append(_topic(subject, sub))

    sheet = {
        "id": "sheet-1",
        "class": "sheet",
        "title": title or "测试用例",
        "rootTopic": _topic(title or "测试用例", root_children),
    }
    content = json.dumps([sheet], ensure_ascii=False, indent=2)

    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
        zf.writestr("content.json", content)
        zf.writestr("metadata.json", json.dumps({"creator": "TestHub"}, ensure_ascii=False))
    return buf.getvalue()


def export_task_result(final_text: str, template_columns, fmt: str, title: str = "测试用例"):
    """统一导出入口。

    参数：
      final_text: 任务最终用例文本（含 Markdown 表格）
      template_columns: 模板列名列表（可空）
      fmt: 'excel' | 'feishu' | 'xmind'
    返回：(content_bytes_or_str, content_type, filename)
    """
    headers, rows = parse_markdown_table(final_text)
    if not headers:
        # 没有表格时，退化为把全文放到 Excel 第一列 / 文本
        rows = [[ln] for ln in (final_text or "").splitlines() if ln.strip()]
        headers = ["内容"]
    out_headers, out_rows = map_to_template(headers, rows, template_columns or [])

    fmt = (fmt or "excel").lower()
    if fmt == "excel":
        data = build_excel_bytes(out_headers, out_rows)
        return data, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "testcases.xlsx"
    if fmt == "xmind":
        data = build_xmind_bytes(out_headers, out_rows, title)
        return data, "application/octet-stream", "testcases.xmind"
    # 默认飞书思维导图
    data = build_feishu_markdown(out_headers, out_rows, title)
    return data, "text/markdown; charset=utf-8", "testcases.md"
