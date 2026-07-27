import logging
import asyncio
import threading
from typing import Any, Dict

from asgiref.sync import async_to_sync
from django.db import close_old_connections
from django.db.models import QuerySet
from django.utils import timezone
from django.http import StreamingHttpResponse, HttpResponse
from rest_framework import status, viewsets
from rest_framework.decorators import action, api_view, permission_classes
from rest_framework.permissions import IsAuthenticated
from rest_framework.response import Response

from .models import (
    AIModelConfig,
    AIModelService,
    AnalysisTask,
    BusinessRequirement,
    GeneratedTestCase,
    GenerationConfig,
    PromptConfig,
    RequirementAnalysis,
    RequirementDocument,
    TestCaseGenerationTask,
    TestCaseSkill,
)
from .serializers import (
    AIModelConfigSerializer,
    AnalysisTaskSerializer,
    BusinessRequirementSerializer,
    DocumentUploadSerializer,
    GeneratedTestCaseSerializer,
    GenerationConfigSerializer,
    PromptConfigSerializer,
    RequirementAnalysisSerializer,
    RequirementDocumentSerializer,
    TestCaseGenerationRequestSerializer,
    TestCaseGenerationCreateFromTextSerializer,
    TestCaseGenerationTaskSerializer,
    TestCaseSkillSerializer,
)
from .services import AIService, RequirementAnalysisService
from .dify_kb_service import (
    fetch_kb_context_for_generation,
    list_datasets,
    list_dataset_documents,
    build_kb_reference_context,
    resolve_dify_config,
)
from .testcase_markdown_utils import remove_cases_by_indices
from .kb_serializers import KbReferencePreviewSerializer

logger = logging.getLogger(__name__)


class RequirementDocumentViewSet(viewsets.ModelViewSet):
    queryset = RequirementDocument.objects.all()
    serializer_class = RequirementDocumentSerializer
    permission_classes = [IsAuthenticated]

    def get_queryset(self) -> QuerySet:
        qs = super().get_queryset()
        project_id = self.request.query_params.get("project")
        if project_id:
            qs = qs.filter(project_id=project_id)
        return qs
    
    def get_serializer_class(self):
        if self.action == "create":
            return DocumentUploadSerializer
        return RequirementDocumentSerializer
    
    @action(detail=True, methods=["post"], url_path="analyze")
    def analyze(self, request, pk=None):
        """对已上传文档触发分析（同步等待完成，适合开发/小文档）。"""
        doc = self.get_object()
        try:
            analysis = async_to_sync(RequirementAnalysisService.process_document_analysis)(doc)
            return Response(RequirementAnalysisSerializer(analysis).data, status=status.HTTP_200_OK)
        except Exception as e:
            logger.exception("文档分析失败")
            return Response({"detail": str(e)}, status=status.HTTP_500_INTERNAL_SERVER_ERROR)

    @action(detail=True, methods=["get"], url_path="extract_text")
    def extract_text(self, request, pk=None):
        """
        兼容旧前端：提取文档文本（/documents/{id}/extract_text/）。
        同时返回 docx 内嵌图片的 data URLs（如有），用于后续“图文一起生成/分析”。
        """
        doc = self.get_object()
        try:
            if not doc.extracted_text:
                doc.extracted_text = RequirementAnalysisService.DocumentProcessor.extract_text(doc)  # type: ignore[attr-defined]
        except Exception:
            # 兼容当前文件组织：DocumentProcessor 在 services.py 中
            from .services import DocumentProcessor

            if not doc.extracted_text:
                doc.extracted_text = DocumentProcessor.extract_text(doc)

        if doc.extracted_text and doc.extracted_text != (doc.extracted_text or ""):
            doc.save(update_fields=["extracted_text"])
        elif not doc.extracted_text:
            doc.save(update_fields=["extracted_text"])

        try:
            from .services import DocumentProcessor

            image_data_urls = DocumentProcessor.extract_image_data_urls(doc)
        except Exception:
            image_data_urls = []

        return Response(
            {
                "id": doc.id,
                "title": doc.title,
                "document_type": doc.document_type,
                "extracted_text": doc.extracted_text or "",
                "image_data_urls": image_data_urls,
                "image_count": len(image_data_urls),
            },
            status=status.HTTP_200_OK,
        )


class RequirementAnalysisViewSet(viewsets.ReadOnlyModelViewSet):
    queryset = RequirementAnalysis.objects.select_related("document").all()
    serializer_class = RequirementAnalysisSerializer
    permission_classes = [IsAuthenticated]

    def get_queryset(self) -> QuerySet:
        qs = super().get_queryset()
        document_id = self.request.query_params.get("document")
        if document_id:
            qs = qs.filter(document_id=document_id)
        return qs


class BusinessRequirementViewSet(viewsets.ModelViewSet):
    queryset = BusinessRequirement.objects.select_related("analysis").all()
    serializer_class = BusinessRequirementSerializer
    permission_classes = [IsAuthenticated]

    def get_queryset(self) -> QuerySet:
        qs = super().get_queryset()
        analysis_id = self.request.query_params.get("analysis")
        if analysis_id:
            qs = qs.filter(analysis_id=analysis_id)
        return qs


class GeneratedTestCaseViewSet(viewsets.ModelViewSet):
    queryset = GeneratedTestCase.objects.select_related("requirement").all()
    serializer_class = GeneratedTestCaseSerializer
    permission_classes = [IsAuthenticated]

    def get_queryset(self) -> QuerySet:
        qs = super().get_queryset()
        requirement_id = self.request.query_params.get("requirement")
        if requirement_id:
            qs = qs.filter(requirement_id=requirement_id)
        return qs


class AnalysisTaskViewSet(viewsets.ReadOnlyModelViewSet):
    queryset = AnalysisTask.objects.select_related("document").all()
    serializer_class = AnalysisTaskSerializer
    permission_classes = [IsAuthenticated]

    def get_queryset(self) -> QuerySet:
        qs = super().get_queryset()
        document_id = self.request.query_params.get("document")
        if document_id:
            qs = qs.filter(document_id=document_id)
        return qs


class AIModelConfigViewSet(viewsets.ModelViewSet):
    queryset = AIModelConfig.objects.all()
    serializer_class = AIModelConfigSerializer
    permission_classes = [IsAuthenticated]

    @action(detail=True, methods=["post"], url_path="test_connection")
    def test_connection(self, request, pk=None):
        """
        前端「测试连接」按钮使用：/requirement-analysis/ai-models/{id}/test_connection/

        说明：
        - 这里只做最小化连通性验证（能否成功请求模型并返回内容）
        - 若配置缺失/网关返回错误，会返回 success=false 并带上错误信息
        """
        cfg: AIModelConfig = self.get_object()
        try:
            messages = [
                {"role": "system", "content": "你是一个测试连接助手。"},
                {"role": "user", "content": "回复 OK（只回复 OK）。"},
            ]
            resp = async_to_sync(AIModelService.call_openai_compatible_api)(cfg, messages)
            content = ""
            try:
                choice0 = (resp.get("choices") or [{}])[0] or {}
                msg = choice0.get("message") or {}
                content = (msg.get("content") or "").strip()
            except Exception:
                content = ""
            return Response(
                {
                    "success": True,
                    "message": "连接成功",
                    "response": content or "OK",
                },
                status=status.HTTP_200_OK,
            )
        except Exception as e:
            return Response(
                {
                    "success": False,
                    "message": str(e) or "连接失败",
                    "response": "",
                },
                status=status.HTTP_200_OK,
            )


class PromptConfigViewSet(viewsets.ModelViewSet):
    queryset = PromptConfig.objects.all()
    serializer_class = PromptConfigSerializer
    permission_classes = [IsAuthenticated]


class GenerationConfigViewSet(viewsets.ModelViewSet):
    queryset = GenerationConfig.objects.all()
    serializer_class = GenerationConfigSerializer
    permission_classes = [IsAuthenticated]


class TestCaseGenerationTaskViewSet(viewsets.ModelViewSet):
    queryset = TestCaseGenerationTask.objects.select_related(
        "project",
        "source_document",
        "writer_model_config",
        "reviewer_model_config",
        "writer_prompt_config",
        "reviewer_prompt_config",
    ).all()
    serializer_class = TestCaseGenerationTaskSerializer
    permission_classes = [IsAuthenticated]
    # 兼容旧前端：按 task_id（TASK_xxx）访问详情
    lookup_field = "task_id"
    lookup_url_kwarg = "task_id"

    def get_object(self):
        """
        兼容两种访问方式：
        - /testcase-generation/TASK_XXXX/  （旧前端使用 task_id）
        - /testcase-generation/123/        （直接用数字主键，便于调试/兼容）
        """
        lookup_value = self.kwargs.get(self.lookup_url_kwarg) or self.kwargs.get("pk")
        queryset = self.filter_queryset(self.get_queryset())
        if isinstance(lookup_value, str) and lookup_value.isdigit():
            obj = queryset.get(pk=int(lookup_value))
        else:
            obj = queryset.get(task_id=lookup_value)
        self.check_object_permissions(self.request, obj)
        return obj

    @action(detail=False, methods=["post"], url_path="create-from-text")
    def create_from_text(self, request):
        """
        兼容旧版：根据纯文本需求创建“生成任务”记录（不立即调用模型）。
        """
        ser = TestCaseGenerationCreateFromTextSerializer(data=request.data)
        ser.is_valid(raise_exception=True)
        task = TestCaseGenerationTask.objects.create(
            title=ser.validated_data["title"],
            requirement_text=ser.validated_data["requirement_text"],
            output_mode=ser.validated_data.get("output_mode", "stream"),
            project_id=ser.validated_data.get("project"),
            source_document_id=ser.validated_data.get("source_document"),
            created_by=request.user,
        )
        return Response(TestCaseGenerationTaskSerializer(task).data, status=status.HTTP_201_CREATED)

    @action(detail=True, methods=["get"], url_path="export")
    def export(self, request, task_id=None):
        """导出生成结果：?format=excel|feishu|xmind（默认 excel）。

        若有团队模板列（skill_template_columns），输出按模板列顺序并做智能字段映射。
        """
        task = self.get_object()
        fmt = (request.query_params.get("format") or "excel").lower()
        if fmt not in ("excel", "feishu", "xmind"):
            fmt = "excel"
        final_text = task.final_test_cases or task.generated_test_cases or ""
        if not final_text.strip():
            return Response({"detail": "该任务还没有可导出的用例内容"}, status=status.HTTP_400_BAD_REQUEST)
        try:
            from .skill_export import export_task_result

            data, content_type, filename = export_task_result(
                final_text=final_text,
                template_columns=task.skill_template_columns or [],
                fmt=fmt,
                title=task.title or "测试用例",
            )
        except Exception as exc:
            logger.exception("导出失败")
            return Response({"detail": f"导出失败：{exc}"}, status=status.HTTP_500_INTERNAL_ERROR)

        resp = HttpResponse(data, content_type=content_type)
        from django.utils.encoding import escape_uri_path

        disp = f"attachment; filename*=UTF-8''{escape_uri_path(filename)}"
        resp["Content-Disposition"] = disp
        return resp

    @action(detail=False, methods=["post"], url_path="generate")
    def generate(self, request):
        """
        兼容旧前端：启动生成任务（/testcase-generation/generate/）。

        旧页面期望返回 {task_id}，随后通过：
        - /testcase-generation/{task_id}/progress/ 轮询
        - /testcase-generation/{task_id}/stream_progress/ SSE（失败时会降级轮询）
        """
        title = (request.data or {}).get("title") or ""
        requirement_text = (request.data or {}).get("requirement_text") or ""
        output_mode = (request.data or {}).get("output_mode") or "stream"
        project_id = (request.data or {}).get("project") or None
        source_document_id = (request.data or {}).get("source_document") or None
        dify_config_id = (request.data or {}).get("dify_config_id") or None
        dify_dataset_id = (request.data or {}).get("dify_dataset_id") or ""
        dify_dataset_name = (request.data or {}).get("dify_dataset_name") or ""
        kb_top_k = (request.data or {}).get("kb_top_k") or 5
        kb_reference_mode = (request.data or {}).get("kb_reference_mode") or "hybrid"
        kb_document_ids = (request.data or {}).get("kb_document_ids") or []
        kb_function_ids = (request.data or {}).get("kb_function_ids") or []
        document_names = (request.data or {}).get("document_names") or {}
        image_data_urls = (request.data or {}).get("image_data_urls") or []
        image_attachments = (request.data or {}).get("image_attachments") or []

        # ── Swagger / OpenAPI 接口文档输入 ──
        swagger_data = (request.data or {}).get("swagger_data")
        if isinstance(swagger_data, str):
            try:
                swagger_data = json.loads(swagger_data)
            except Exception:
                swagger_data = None
        # 也支持直接把 OpenAPI JSON 作为 requirement_text 传入（自动识别）
        _maybe_swagger = None
        if not swagger_data and requirement_text.strip().startswith("{"):
            try:
                _cand = json.loads(requirement_text)
                if isinstance(_cand, dict) and ("openapi" in _cand or "swagger" in _cand or "paths" in _cand):
                    swagger_data = _cand
            except Exception:
                pass
        if swagger_data and isinstance(swagger_data, dict):
            try:
                swagger_text = _swagger_to_requirement_text(swagger_data)
                requirement_text = (
                    f"【接口文档（Swagger/OpenAPI）】请据此生成接口测试用例，覆盖每个接口的正常/异常/边界场景。\n\n"
                    f"{swagger_text}\n\n"
                    f"{requirement_text}"
                ).strip()
            except Exception as exc:
                logger.warning("Swagger 解析失败: %s", exc)

        if not isinstance(kb_document_ids, list):
            kb_document_ids = []
        if not isinstance(kb_function_ids, list):
            kb_function_ids = []
        if not isinstance(image_data_urls, list):
            image_data_urls = []
        if not isinstance(image_attachments, list):
            image_attachments = []

        from .image_attachment_utils import normalize_image_attachments

        normalized_images = normalize_image_attachments(image_attachments, image_data_urls)
        image_data_urls = [a["url"] for a in normalized_images]
        kb_document_ids = [str(x).strip() for x in kb_document_ids if str(x).strip()]
        try:
            kb_function_ids = [int(x) for x in kb_function_ids]
        except (TypeError, ValueError):
            kb_function_ids = []

        if not str(title).strip() or not str(requirement_text).strip():
            return Response({"detail": "title 和 requirement_text 不能为空"}, status=status.HTTP_400_BAD_REQUEST)

        enable_kb = (request.data or {}).get("enable_knowledge_base", False)
        if isinstance(enable_kb, str):
            enable_kb = enable_kb.lower() in ("true", "1", "yes", "on")
        if not isinstance(enable_kb, bool):
            enable_kb = bool(enable_kb)

        kb_context = ""
        kb_context_meta: Dict[str, Any] = {}
        dify_config = None
        has_kb_refs = bool(kb_document_ids) or bool(kb_function_ids)

        # ── 新模式：启用知识库(业务大脑) + 关联项目 → 自动检索项目知识库 ──
        # 支持 Dify 引擎（项目绑定的 Dify 数据集）与 native 引擎（项目已发布自建 KB），
        # 并尊重前端勾选的 selected_kb_ids（未勾选时取全部可用 KB）。
        if enable_kb and project_id and not has_kb_refs:
            try:
                from .kb_hub.backend import KbBackendFactory
                from .kb_hub.models import NativeKb, ProjectDifyKbBinding
                from apps.projects.models import Project

                selected_kb_ids = [
                    str(x).strip() for x in (request.data or {}).get("selected_kb_ids") or []
                    if str(x).strip()
                ]

                proj = Project.objects.filter(pk=project_id).first()
                if proj:
                    engine = KbBackendFactory.get_engine()
                    try:
                        kb_top_k = int(kb_top_k)
                    except (TypeError, ValueError):
                        kb_top_k = 5
                    kb_top_k = max(1, min(kb_top_k, 10))

                    # 收集 (kb_id, kb_name) 列表：Dify 引擎用 dataset_id，native 用 NativeKb.pk
                    entries: list = []
                    if engine == "dify":
                        bindings = ProjectDifyKbBinding.objects.filter(project_id=project_id)
                        if selected_kb_ids:
                            bindings = bindings.filter(dataset_id__in=selected_kb_ids)
                        entries = [(b.dataset_id, b.dataset_name) for b in bindings]
                        if entries:
                            # 记录使用的 Dify 配置，供后续阶段/展示使用
                            first_binding = bindings.first()
                            dify_config = resolve_dify_config(first_binding.dify_config_id if first_binding else None)
                    else:
                        native_qs = proj.native_kbs.filter(status="published")
                        if selected_kb_ids:
                            int_ids = [int(x) for x in selected_kb_ids if str(x).isdigit()]
                            if int_ids:
                                native_qs = native_qs.filter(pk__in=int_ids)
                        entries = [(str(nkb.pk), nkb.name) for nkb in native_qs]

                    if entries:
                        all_sources: list = []
                        all_context_parts: list = []
                        kb_names: list = []
                        for kb_id, kb_name in entries:
                            try:
                                backend = KbBackendFactory.get_backend()
                                ctx, meta = backend.build_generation_context(
                                    title=str(title).strip(),
                                    requirement_text=str(requirement_text).strip(),
                                    kb_name=kb_name,
                                    document_ids=[],
                                    function_ids=[],
                                    reference_mode="full",
                                    top_k=kb_top_k,
                                    kb_id=str(kb_id),
                                )
                                if ctx.strip():
                                    all_context_parts.append(ctx)
                                srcs = meta.get("sources") or []
                                for s in srcs:
                                    s.setdefault("kb_name", kb_name)
                                all_sources.extend(srcs)
                                kb_names.append(kb_name)
                            except Exception as exc:
                                logger.warning("知识库 %s 检索失败: %s", kb_name, exc, exc_info=True)

                        kb_context = "\n\n".join(all_context_parts)
                        kb_context_meta = {
                            "engine": engine,
                            "mode": "project_auto_retrieve",
                            "kb_count": len(entries),
                            "kb_names": kb_names,
                            "sources": all_sources,
                            "total_chars": len(kb_context),
                        }
                        if engine == "dify":
                            dify_dataset_id = entries[0][0]
                            dify_dataset_name = "、".join(kb_names)
                        else:
                            dify_dataset_name = "、".join(kb_names)
                    else:
                        logger.info("项目 %s 没有可用的知识库（引擎=%s），跳过自动检索", proj.name, engine)
            except Exception as exc:
                logger.warning("启用知识库自动检索失败: %s", exc, exc_info=True)

        # ── 旧模式：手动选择 Dify 数据集 + 文档/功能模块 ──
        elif str(dify_dataset_id).strip() and has_kb_refs:
            from .kb_hub.backend import KbBackendFactory
            engine = KbBackendFactory.get_engine()
            try:
                kb_top_k = int(kb_top_k)
            except (TypeError, ValueError):
                kb_top_k = 5
            kb_top_k = max(1, min(kb_top_k, 10))
            ref_mode = "documents"
            dataset_id_str = str(dify_dataset_id).strip()
            doc_names_map = document_names if isinstance(document_names, dict) else {}
            expansion: Dict[str, Any] = {}
            # Dify 引擎下解析 dify_config 并执行图谱扩展；native 引擎跳过
            if engine == "dify":
                dify_config = resolve_dify_config(dify_config_id)
                if not dify_config:
                    return Response({"detail": "未找到可用的 Dify 配置，请先在配置中心添加并启用，或切换知识中枢引擎"}, status=status.HTTP_400_BAD_REQUEST)
                try:
                    from apps.knowledge_graph.query import expand_kb_references_via_graph

                    expansion = expand_kb_references_via_graph(
                        dataset_id=dataset_id_str,
                        function_ids=kb_function_ids,
                        document_ids=kb_document_ids,
                        max_depth=2,
                    )
                    kb_function_ids = expansion.get("function_ids") or kb_function_ids
                    kb_document_ids = expansion.get("document_ids") or kb_document_ids
                    doc_names_map = {**doc_names_map, **(expansion.get("document_names") or {})}
                except Exception as exc:
                    logger.warning("图谱扩展 KB 参考失败，使用原始选择: %s", exc, exc_info=True)
            try:
                backend = KbBackendFactory.get_backend()
                kb_context, kb_context_meta = backend.build_generation_context(
                    title=str(title).strip(),
                    requirement_text=str(requirement_text).strip(),
                    kb_name=str(dify_dataset_name).strip(),
                    document_ids=kb_document_ids,
                    function_ids=kb_function_ids,
                    reference_mode=ref_mode,
                    top_k=kb_top_k,
                    kb_id=dataset_id_str,
                    document_names=doc_names_map if engine == "dify" else None,
                )
                if expansion.get("expanded"):
                    kb_context_meta = kb_context_meta or {}
                    kb_context_meta["graph_expansion"] = expansion.get("meta") or {}
                    if expansion.get("graph_summary"):
                        kb_context_meta["graph_summary"] = expansion["graph_summary"]
                try:
                    from apps.knowledge_graph.query import build_graph_relations_prompt_summary
                    from types import SimpleNamespace

                    _shim = SimpleNamespace(
                        kb_function_ids=kb_function_ids,
                        kb_document_ids=kb_document_ids,
                        kb_context_meta=kb_context_meta or {},
                    )
                    _prompt_summary = build_graph_relations_prompt_summary(_shim)
                    if _prompt_summary:
                        kb_context_meta = kb_context_meta or {}
                        kb_context_meta["graph_prompt_injected"] = True
                        kb_context_meta["graph_prompt_summary_chars"] = len(_prompt_summary)
                except Exception as exc:
                    logger.warning("图谱 prompt 摘要统计失败: %s", exc, exc_info=True)
            except ValueError as exc:
                return Response({"detail": str(exc)}, status=status.HTTP_400_BAD_REQUEST)
            except Exception as exc:
                logger.exception("Dify 知识库参考构建失败")
                return Response({"detail": f"知识库参考构建失败: {exc}"}, status=status.HTTP_502_BAD_GATEWAY)
            if not kb_context.strip():
                return Response(
                    {"detail": "未能获取有效的知识库参考内容，请检查文档选择"},
                    status=status.HTTP_400_BAD_REQUEST,
                )

        # 选取当前启用的配置（与旧页面的 config/check 逻辑保持一致）
        # 如果传了 skill_id，优先使用 Skill 中绑定的配置
        skill_id = (request.data or {}).get("skill_id")
        skill = None
        if skill_id:
            try:
                skill = TestCaseSkill.objects.get(pk=skill_id, is_active=True)
            except TestCaseSkill.DoesNotExist:
                return Response({"detail": "指定的 Skill 不存在或未启用"}, status=status.HTTP_400_BAD_REQUEST)

        if skill:
            writer_model = skill.writer_model_config
            reviewer_model = skill.reviewer_model_config
            writer_prompt = skill.writer_prompt_config
            reviewer_prompt = skill.reviewer_prompt_config
            gen_cfg = skill.generation_config or GenerationConfig.get_active_config()
        else:
            writer_model = AIModelConfig.objects.filter(role="writer", is_active=True).order_by("-updated_at").first()
            reviewer_model = AIModelConfig.objects.filter(role="reviewer", is_active=True).order_by("-updated_at").first()
            writer_prompt = PromptConfig.objects.filter(prompt_type="writer", is_active=True).order_by("-updated_at").first()
            reviewer_prompt = PromptConfig.objects.filter(prompt_type="reviewer", is_active=True).order_by("-updated_at").first()
            gen_cfg = GenerationConfig.get_active_config()

        # 组装 Skill 注入提示词（系统提示词 + 约束规则 + 模板列要求）
        skill_system_prompt = ""
        skill_template_columns = []
        if skill:
            parts = [skill.get_full_system_prompt()]
            tcols = skill.template_columns or []
            if tcols:
                skill_template_columns = list(tcols)
                cols_text = "、".join(str(c) for c in tcols)
                parts.append(
                    f"\n\n# 输出列要求（团队模板）\n"
                    f"你必须严格按照以下列名和顺序输出测试用例表格，不要增删或改名：\n"
                    f"{cols_text}\n"
                    f"每一行是一条用例，缺失信息留空即可。输出格式为 Markdown 表格。"
                )
            skill_system_prompt = "\n".join(parts)

        if not (writer_model and writer_prompt):
            return Response({"detail": "未配置启用的编写模型/提示词"}, status=status.HTTP_400_BAD_REQUEST)

        enable_auto_review = True
        if gen_cfg and gen_cfg.enable_auto_review is not None:
            enable_auto_review = bool(gen_cfg.enable_auto_review)
        review_timeout_s = None
        if gen_cfg and getattr(gen_cfg, "review_timeout", None):
            try:
                review_timeout_s = int(gen_cfg.review_timeout)
            except Exception:
                review_timeout_s = None

        # 任务号：TASK_xxxxxxxx
        import uuid

        task = TestCaseGenerationTask.objects.create(
            task_id=f"TASK_{uuid.uuid4().hex[:8].upper()}",
            title=str(title).strip(),
            requirement_text=str(requirement_text).strip(),
            output_mode="stream" if output_mode == "stream" else "complete",
            project_id=project_id or None,
            source_document_id=source_document_id or None,
            dify_config=dify_config,
            dify_dataset_id=str(dify_dataset_id).strip() if str(dify_dataset_id).strip() else "",
            dify_dataset_name=str(dify_dataset_name).strip(),
            kb_context=kb_context or "",
            kb_context_meta=kb_context_meta or {},
            kb_top_k=kb_top_k if str(dify_dataset_id).strip() else 5,
            kb_reference_mode="documents" if has_kb_refs else "hybrid",
            kb_document_ids=kb_document_ids if has_kb_refs else [],
            kb_function_ids=kb_function_ids if has_kb_refs else [],
            image_data_urls=image_data_urls,
            image_attachments=normalized_images,
            writer_model_config=writer_model,
            reviewer_model_config=reviewer_model,
            writer_prompt_config=writer_prompt,
            reviewer_prompt_config=reviewer_prompt,
            skill=skill,
            skill_system_prompt=skill_system_prompt,
            skill_template_columns=skill_template_columns,
            created_by=request.user,
            status="pending",
            progress=0,
        )

        # 后台执行：避免阻塞接口响应（开发环境可用；生产建议 Celery）
        from .generation_task_runner import start_generation_task_background

        start_generation_task_background(task.pk)

        kb_chat_session_id = str((request.data or {}).get("kb_chat_session_id") or "").strip()
        if kb_chat_session_id:
            from django.utils import timezone as tz

            from .kb_chat_models import KbChatSession

            KbChatSession.objects.filter(
                session_id=kb_chat_session_id,
                user=request.user,
            ).update(
                last_generation_task_id=task.task_id,
                updated_at=tz.now(),
            )

        try:
            from apps.knowledge_graph.builder import index_generation_task

            index_generation_task(
                task,
                created_by=request.user,
                kb_chat_session_id=kb_chat_session_id or None,
            )
        except Exception:
            pass

        return Response({"task_id": task.task_id}, status=status.HTTP_201_CREATED)

    @action(detail=False, methods=["post"], url_path="upload-images")
    def upload_images(self, request):
        """上传界面截图，返回多模态 data URLs（供 generate 接口使用）。"""
        from .services import DocumentProcessor

        files = request.FILES.getlist("files") or request.FILES.getlist("file") or []
        if not files:
            return Response({"detail": "请上传至少一张图片"}, status=status.HTTP_400_BAD_REQUEST)
        urls = DocumentProcessor.upload_files_to_data_urls(files)
        if not urls:
            return Response({"detail": "无法解析图片，请上传 PNG/JPG/WebP 格式"}, status=status.HTTP_400_BAD_REQUEST)
        attachments = [{"url": u, "role": "ui_layout"} for u in urls]
        return Response(
            {"image_data_urls": urls, "image_attachments": attachments, "count": len(urls)},
            status=status.HTTP_200_OK,
        )

    @action(detail=True, methods=["post"], url_path="continue-refine")
    def continue_refine(self, request, task_id=None):
        """
        在已有用例基础上继续优化：补充文字要求、上传新截图（页面样式/操作步骤）。
        复用同一 task_id，完成后刷新用例列表。
        """
        task = self.get_object()
        if task.status in ("pending", "generating", "reviewing", "revising"):
            return Response({"detail": "任务进行中，请稍后再试"}, status=status.HTTP_400_BAD_REQUEST)

        existing_cases = (task.final_test_cases or task.generated_test_cases or "").strip()
        if not existing_cases:
            return Response(
                {"detail": "当前任务尚无可用用例，无法在此基础上继续修改"},
                status=status.HTTP_400_BAD_REQUEST,
            )

        refinement_instructions = str((request.data or {}).get("refinement_instructions") or "").strip()
        new_attachments = (request.data or {}).get("image_attachments") or []
        new_urls = (request.data or {}).get("image_data_urls") or []
        if not isinstance(new_attachments, list):
            new_attachments = []
        if not isinstance(new_urls, list):
            new_urls = []

        from .image_attachment_utils import (
            merge_image_attachments,
            normalize_image_attachments,
        )

        existing_att = normalize_image_attachments(task.image_attachments, task.image_data_urls)
        incoming_att = normalize_image_attachments(new_attachments, new_urls)
        merged_att = merge_image_attachments(existing_att, incoming_att)

        if not refinement_instructions and not incoming_att:
            return Response(
                {"detail": "请填写补充要求，或上传新的参考截图"},
                status=status.HTTP_400_BAD_REQUEST,
            )

        if not task.writer_model_config_id or not task.writer_prompt_config_id:
            return Response({"detail": "任务缺少编写模型/提示词配置"}, status=status.HTTP_400_BAD_REQUEST)

        stamp = timezone.now().strftime("%Y-%m-%d %H:%M")
        note_block = f"\n\n--- 迭代补充 {stamp} ---\n"
        if refinement_instructions:
            note_block += refinement_instructions
        if incoming_att:
            note_block += f"\n（新增截图 {len(incoming_att)} 张）"
        refinement_notes = ((task.refinement_notes or "").strip() + note_block).strip()

        task.refinement_notes = refinement_notes
        task.image_attachments = merged_att
        task.image_data_urls = [a["url"] for a in merged_att]
        task.stream_buffer = ""
        task.stream_position = 0
        task.status = "pending"
        task.progress = 0
        task.error_message = ""
        task.completed_at = None
        task.save(
            update_fields=[
                "refinement_notes",
                "image_attachments",
                "image_data_urls",
                "stream_buffer",
                "stream_position",
                "status",
                "progress",
                "error_message",
                "completed_at",
            ]
        )

        from .generation_task_runner import start_refinement_task_background

        start_refinement_task_background(task.pk, refinement_instructions)

        try:
            from apps.knowledge_graph.builder import index_generation_task

            index_generation_task(task, created_by=request.user)
        except Exception:
            pass

        return Response(
            {
                "task_id": task.task_id,
                "detail": "已开始基于现有用例继续优化",
                "image_count": len(merged_att),
            },
            status=status.HTTP_200_OK,
        )

    @action(detail=True, methods=["post"], url_path="save_to_records")
    def save_to_records(self, request, task_id=None):
        """
        一键采纳：将任务的最终测试用例导入到“测试用例”模块（apps.testcases.TestCase）。

        约束：
        - 任务必须有关联 project，否则无法创建 TestCase（TestCase.project 为必填）
        - 优先解析 JSON（数组/对象）；解析失败则做尽力的文本分段解析
        """
        task = self.get_object()
        if task.is_saved_to_records:
            return Response({"created_count": 0, "detail": "已采纳"}, status=status.HTTP_200_OK)

        if not task.project_id:
            return Response({"detail": "该任务未关联项目，无法一键采纳（请在生成时选择项目）"}, status=status.HTTP_400_BAD_REQUEST)

        raw = (task.final_test_cases or task.generated_test_cases or "").strip()
        if not raw:
            return Response({"detail": "任务无可采纳的用例内容"}, status=status.HTTP_400_BAD_REQUEST)

        import json
        import re
        from apps.testcases.models import TestCase

        def _strip_fence(s: str) -> str:
            s = (s or "").strip()
            s = re.sub(r"^```[a-zA-Z0-9_-]*\s*", "", s)
            s = re.sub(r"\s*```$", "", s)
            return s.strip()

        def _try_parse_json(s: str):
            s2 = _strip_fence(s)
            try:
                obj = json.loads(s2)
            except Exception:
                return None
            if isinstance(obj, dict):
                for k in ("testcases", "cases", "data", "items"):
                    v = obj.get(k)
                    if isinstance(v, list):
                        return v
                return [obj]
            if isinstance(obj, list):
                return obj
            return None

        def _pick(d: dict, *keys, default=""):
            for k in keys:
                if k in d and d[k] not in (None, ""):
                    return d[k]
            return default

        def _normalize_priority(v: str) -> str:
            v = (v or "").lower().strip()
            if v in ("low", "medium", "high", "critical"):
                return v
            mp = {"p0": "critical", "p1": "high", "p2": "medium", "p3": "low"}
            return mp.get(v, "medium")

        def _normalize_type(v: str) -> str:
            v = (v or "").lower().strip()
            allowed = {"functional", "integration", "api", "ui", "performance", "security"}
            return v if v in allowed else "functional"

        cases = []
        json_cases = _try_parse_json(raw)
        if json_cases:
            for it in json_cases:
                if isinstance(it, dict):
                    cases.append(it)

        # 兜底：按空行分段，每段视为一个用例（保守策略，避免完全不可用）
        if not cases:
            blocks = [b.strip() for b in raw.split("\n\n") if b.strip()]
            for b in blocks[:200]:
                lines = [ln.strip() for ln in b.splitlines() if ln.strip()]
                title = lines[0][:200] if lines else task.title
                cases.append(
                    {
                        "title": title,
                        "preconditions": "",
                        "steps": b[:1000],
                        "expected_result": "见步骤/描述",
                        "priority": "medium",
                        "status": "active",
                        "test_type": "functional",
                        "description": "",
                    }
                )

        created = 0
        for d in cases[:500]:
            title = str(_pick(d, "title", "name", "用例标题", default=task.title)).strip()[:500] or task.title
            preconditions = str(_pick(d, "preconditions", "precondition", "前置条件", default="")).strip()
            steps = str(_pick(d, "steps", "test_steps", "步骤", "操作步骤", default="")).strip()
            expected = str(_pick(d, "expected_result", "expected", "预期结果", default="见步骤/描述")).strip() or "见步骤/描述"
            description = str(_pick(d, "description", "desc", "用例描述", default="")).strip()
            priority = _normalize_priority(str(_pick(d, "priority", "优先级", default="medium")))
            status_v = str(_pick(d, "status", default="active")).strip() or "active"
            if status_v not in ("draft", "active", "deprecated"):
                status_v = "active"
            test_type = _normalize_type(str(_pick(d, "test_type", "type", "测试类型", default="functional")))

            tc = TestCase.objects.create(
                project_id=task.project_id,
                title=title,
                description=description,
                preconditions=preconditions,
                steps=steps[:1000],
                expected_result=expected,
                priority=priority,
                status=status_v,
                test_type=test_type,
                tags=[],
                author=task.created_by,
            )
            try:
                from apps.knowledge_graph.builder import index_adopted_test_case

                index_adopted_test_case(tc, task, created_by=request.user)
            except Exception:
                pass
            created += 1

        task.is_saved_to_records = True
        task.saved_at = timezone.now()
        task.save(update_fields=["is_saved_to_records", "saved_at"])
        return Response({"created_count": created}, status=status.HTTP_200_OK)

    @action(detail=True, methods=["post"], url_path="discard")
    def discard(self, request, task_id=None):
        """一键弃用：将任务标记为已取消（保留记录，便于追溯）。"""
        task = self.get_object()
        if task.status != "completed":
            return Response({"detail": "仅已完成的任务才支持一键弃用"}, status=status.HTTP_400_BAD_REQUEST)
        task.status = "cancelled"
        task.save(update_fields=["status"])
        return Response({"detail": "ok"}, status=status.HTTP_200_OK)

    @action(detail=True, methods=["post"], url_path="batch-adopt-selected")
    def batch_adopt_selected(self, request, task_id=None):
        """批量采纳选中的测试用例到测试用例模块。"""
        from apps.projects.models import Project
        from apps.testcases.models import TestCase

        task = self.get_object()
        items = request.data.get("test_cases") or []
        if not isinstance(items, list) or not items:
            return Response({"detail": "test_cases 不能为空"}, status=status.HTTP_400_BAD_REQUEST)

        project = task.project
        if not project:
            project_id = request.data.get("project") or request.data.get("project_id")
            if project_id:
                project = Project.objects.filter(pk=project_id).first()
            if not project:
                project = Project.objects.first()
            if not project:
                project = Project.objects.create(
                    name="默认项目",
                    owner=request.user,
                    description="系统自动创建的默认项目",
                )

        priority_map = {
            "critical": "critical",
            "high": "high",
            "medium": "medium",
            "low": "low",
        }

        created = 0
        author = task.created_by or request.user
        for item in items[:500]:
            if not isinstance(item, dict):
                continue
            priority = str(item.get("priority") or "medium").lower()
            if priority not in priority_map:
                priority = "medium"
            TestCase.objects.create(
                project=project,
                title=str(item.get("title") or item.get("scenario") or "测试用例")[:500],
                description=str(item.get("description") or item.get("scenario") or "")[:2000],
                preconditions=str(item.get("preconditions") or item.get("precondition") or "")[:2000],
                steps=str(item.get("steps") or "")[:1000],
                expected_result=str(item.get("expected_result") or item.get("expected") or "见步骤")[:2000],
                priority=priority,
                status=str(item.get("status") or "draft"),
                test_type=str(item.get("test_type") or "functional"),
                tags=[],
                author=author,
            )
            created += 1

        return Response({"created_count": created, "message": "ok"}, status=status.HTTP_200_OK)

    @action(detail=True, methods=["post"], url_path="discard-selected-cases")
    def discard_selected_cases(self, request, task_id=None):
        """从任务中移除选中的测试用例行。"""
        task = self.get_object()
        indices = request.data.get("case_indices") or []
        if not isinstance(indices, list) or not indices:
            return Response({"detail": "case_indices 不能为空"}, status=status.HTTP_400_BAD_REQUEST)

        raw = (task.final_test_cases or task.generated_test_cases or "").strip()
        if not raw:
            return Response({"detail": "任务无用例内容"}, status=status.HTTP_400_BAD_REQUEST)

        new_raw, removed, remaining = remove_cases_by_indices(raw, indices)
        if not remaining:
            task.final_test_cases = ""
            task.generated_test_cases = ""
            task.status = "cancelled"
            task.save(update_fields=["final_test_cases", "generated_test_cases", "status"])
            return Response(
                {"task_deleted": True, "discarded_count": removed, "updated_test_cases": ""},
                status=status.HTTP_200_OK,
            )

        task.final_test_cases = new_raw
        task.generated_test_cases = new_raw
        task.save(update_fields=["final_test_cases", "generated_test_cases"])
        return Response(
            {
                "task_deleted": False,
                "discarded_count": removed,
                "updated_test_cases": new_raw,
            },
            status=status.HTTP_200_OK,
        )

    @action(detail=True, methods=["post"], url_path="discard-single-case")
    def discard_single_case(self, request, task_id=None):
        """弃用任务中的单条测试用例。"""
        case_index = request.data.get("case_index")
        if case_index is None:
            return Response({"detail": "case_index 不能为空"}, status=status.HTTP_400_BAD_REQUEST)

        task = self.get_object()
        raw = (task.final_test_cases or task.generated_test_cases or "").strip()
        if not raw:
            return Response({"detail": "任务无用例内容"}, status=status.HTTP_400_BAD_REQUEST)
        new_raw, removed, remaining = remove_cases_by_indices(raw, [case_index])
        if not remaining:
            task.final_test_cases = ""
            task.generated_test_cases = ""
            task.status = "cancelled"
            task.save(update_fields=["final_test_cases", "generated_test_cases", "status"])
            return Response(
                {"task_deleted": True, "discarded_count": removed, "updated_test_cases": ""},
                status=status.HTTP_200_OK,
            )
        task.final_test_cases = new_raw
        task.generated_test_cases = new_raw
        task.save(update_fields=["final_test_cases", "generated_test_cases"])
        return Response(
            {
                "task_deleted": False,
                "discarded_count": removed,
                "updated_test_cases": new_raw,
            },
            status=status.HTTP_200_OK,
        )

    @action(detail=True, methods=["get"], url_path="progress")
    def progress(self, request, task_id=None):
        """兼容旧前端：获取任务进度（/testcase-generation/{task_id}/progress/）。"""
        task = self.get_object()
        return Response(TestCaseGenerationTaskSerializer(task).data, status=status.HTTP_200_OK)

    @action(detail=True, methods=["get"], url_path="stream_progress")
    def stream_progress(self, request, task_id=None):
        """
        兼容旧前端：SSE 推送进度与内容（/testcase-generation/{task_id}/stream_progress/）。

        说明：
        - 旧前端用 EventSource（无法携带 Bearer 头）；若鉴权失败会自动降级到轮询。
        - 这里按“每行一个 JSON event”从 task.stream_buffer 输出。
        """
        task = self.get_object()

        def event_stream():
            import time as _time

            last_pos = int(getattr(task, "stream_position", 0) or 0)
            idle = 0
            while True:
                # 重新取最新任务状态
                t = TestCaseGenerationTask.objects.get(pk=task.pk)
                buf = t.stream_buffer or ""
                if last_pos < len(buf):
                    chunk = buf[last_pos:]
                    last_pos = len(buf)
                    # 逐行发送，避免拆 JSON
                    for line in chunk.splitlines():
                        line = line.strip()
                        if not line:
                            continue
                        yield f"data: {line}\n\n"
                    # 记录位置
                    t.stream_position = last_pos
                    t.save(update_fields=["stream_position"])
                    idle = 0
                else:
                    idle += 1

                if t.status in ("completed", "failed", "cancelled"):
                    # 等待 buffer 发完后退出
                    if last_pos >= len(t.stream_buffer or ""):
                        break
                # 空闲过久也退出，避免连接永远挂着
                if idle > 300:
                    break
                _time.sleep(1)

        resp = StreamingHttpResponse(event_stream(), content_type="text/event-stream")
        resp["Cache-Control"] = "no-cache"
        return resp


class DifyKnowledgeBaseViewSet(viewsets.ViewSet):
    """Dify 知识库列表（供 AI 用例生成时选择）"""
    permission_classes = [IsAuthenticated]

    def list(self, request):
        dify_config_id = request.query_params.get("dify_config_id")
        keyword = (request.query_params.get("keyword") or "").strip()
        try:
            page = int(request.query_params.get("page") or 1)
        except (TypeError, ValueError):
            page = 1
        try:
            limit = int(request.query_params.get("limit") or 100)
        except (TypeError, ValueError):
            limit = 100

        config = resolve_dify_config(dify_config_id)
        if not config:
            return Response(
                {"detail": "未找到可用的 Dify 配置，请先在配置中心添加并启用"},
                status=status.HTTP_400_BAD_REQUEST,
            )
        try:
            payload = list_datasets(config, page=page, limit=limit, keyword=keyword)
            payload["dify_config_id"] = config.id
            return Response(payload)
        except ValueError as exc:
            # 配置错误 / Dify 不可达（service 层已转为可读中文消息）
            return Response({"detail": str(exc)}, status=status.HTTP_400_BAD_REQUEST)
        except Exception as exc:
            logger.exception("获取 Dify 知识库列表失败")
            return Response(
                {"detail": f"获取知识库列表失败（{type(exc).__name__}）：请稍后重试或联系管理员。"},
                status=status.HTTP_502_BAD_GATEWAY,
            )

    @action(detail=False, methods=["get"], url_path="documents")
    def documents(self, request):
        dify_config_id = request.query_params.get("dify_config_id")
        dataset_id = (request.query_params.get("dataset_id") or "").strip()
        keyword = (request.query_params.get("keyword") or "").strip()
        if not dataset_id:
            return Response({"detail": "dataset_id 不能为空"}, status=status.HTTP_400_BAD_REQUEST)

        config = resolve_dify_config(dify_config_id)
        if not config:
            return Response({"detail": "未找到可用的 Dify 配置"}, status=status.HTTP_400_BAD_REQUEST)

        try:
            page = int(request.query_params.get("page") or 1)
        except (TypeError, ValueError):
            page = 1
        try:
            limit = int(request.query_params.get("limit") or 100)
        except (TypeError, ValueError):
            limit = 100

        try:
            payload = list_dataset_documents(
                config, dataset_id, page=page, limit=limit, keyword=keyword
            )
            payload["dify_config_id"] = config.id
            payload["dataset_id"] = dataset_id
            return Response(payload)
        except ValueError as exc:
            return Response({"detail": str(exc)}, status=status.HTTP_400_BAD_REQUEST)
        except Exception as exc:
            logger.exception("获取知识库文档失败")
            return Response(
                {"detail": f"获取知识库文档失败（{type(exc).__name__}）：请稍后重试。"},
                status=status.HTTP_502_BAD_GATEWAY,
            )

    @action(detail=False, methods=["post"], url_path="preview-reference")
    def preview_reference(self, request):
        ser = KbReferencePreviewSerializer(data=request.data)
        ser.is_valid(raise_exception=True)
        data = ser.validated_data

        doc_ids = data.get("kb_document_ids") or []
        func_ids = data.get("kb_function_ids") or []
        if not doc_ids and not func_ids:
            return Response(
                {
                    "context_preview": "",
                    "context_length": 0,
                    "meta": {"skipped": "未选择参考文档或功能模块，不进行知识库拉取"},
                    "has_content": False,
                }
            )

        config = resolve_dify_config(data.get("dify_config_id"))
        if not config:
            return Response({"detail": "未找到可用的 Dify 配置"}, status=status.HTTP_400_BAD_REQUEST)

        try:
            context, meta = build_kb_reference_context(
                config,
                data["dify_dataset_id"],
                title=data.get("title") or "",
                requirement_text=data.get("requirement_text") or "",
                dataset_name=data.get("dify_dataset_name") or "",
                document_ids=data.get("kb_document_ids") or [],
                document_names=data.get("document_names") or {},
                function_ids=data.get("kb_function_ids") or [],
                reference_mode="documents",
                top_k=data.get("kb_top_k") or 5,
            )
            return Response(
                {
                    "context_preview": context,
                    "context_length": len(context),
                    "meta": meta,
                    "has_content": bool(context.strip()),
                    "is_truncated": False,
                }
            )
        except ValueError as exc:
            return Response({"detail": str(exc)}, status=status.HTTP_400_BAD_REQUEST)
        except Exception as exc:
            logger.exception("预览知识库参考失败")
            return Response({"detail": str(exc)}, status=status.HTTP_502_BAD_GATEWAY)


class TestCaseSkillViewSet(viewsets.ModelViewSet):
    """Skill 技能包配置 ViewSet"""
    queryset = TestCaseSkill.objects.select_related(
        'writer_model_config', 'reviewer_model_config',
        'writer_prompt_config', 'reviewer_prompt_config',
        'generation_config', 'created_by',
    ).all()
    serializer_class = TestCaseSkillSerializer
    permission_classes = [IsAuthenticated]

    def get_queryset(self):
        qs = super().get_queryset()
        skill_type = self.request.query_params.get('skill_type')
        if skill_type:
            qs = qs.filter(skill_type=skill_type)
        return qs

    def perform_destroy(self, instance):
        """内置 Skill 不可删除"""
        if instance.is_builtin:
            from rest_framework.exceptions import PermissionDenied
            raise PermissionDenied("内置 Skill 不可删除，可复制后修改")
        super().perform_destroy(instance)

    @action(detail=False, methods=["get"], url_path="active")
    def active(self, request):
        """返回所有启用的 Skill 列表"""
        skills = TestCaseSkill.get_active_skills()
        ser = TestCaseSkillSerializer(skills, many=True)
        return Response(ser.data, status=status.HTTP_200_OK)

    @action(detail=True, methods=["post"], url_path="duplicate")
    def duplicate(self, request, pk=None):
        """复制 Skill（内置 Skill 也可复制为自定义）"""
        skill = self.get_object()
        import copy

        new_skill = TestCaseSkill(
            name=f"{skill.name} (副本)",
            description=skill.description,
            icon=skill.icon,
            category=skill.category,
            skill_type=skill.skill_type,
            system_prompt=skill.system_prompt,
            constraint_rules=skill.constraint_rules,
            output_format=skill.output_format,
            is_builtin=False,  # 副本永远不是内置
            version='1.0',
            writer_model_config=skill.writer_model_config,
            reviewer_model_config=skill.reviewer_model_config,
            writer_prompt_config=skill.writer_prompt_config,
            reviewer_prompt_config=skill.reviewer_prompt_config,
            generation_config=skill.generation_config,
            is_active=True,
            sort_order=skill.sort_order + 1,
            created_by=request.user if request.user.is_authenticated else skill.created_by,
        )
        new_skill.save()
        ser = TestCaseSkillSerializer(new_skill)
        return Response(ser.data, status=status.HTTP_201_CREATED)

    @action(detail=True, methods=["post"], url_path="upload-template")
    def upload_template(self, request, pk=None):
        """上传团队模板（Excel/CSV），自动解析首行列名。"""
        skill = self.get_object()
        files = request.FILES.getlist("file") or request.FILES.getlist("files")
        if not files:
            return Response({"detail": "请选择模板文件"}, status=status.HTTP_400_BAD_REQUEST)
        f = files[0]
        try:
            columns = _parse_template_columns(f)
        except Exception as exc:
            return Response({"detail": f"模板解析失败：{exc}"}, status=status.HTTP_400_BAD_REQUEST)
        if not columns:
            return Response({"detail": "未能从模板首行解析到列名"}, status=status.HTTP_400_BAD_REQUEST)
        # 删除旧文件，避免堆积
        if skill.template_file:
            try:
                skill.template_file.delete(save=False)
            except Exception:
                pass
        skill.template_file = f
        skill.template_columns = columns
        skill.save()
        ser = TestCaseSkillSerializer(skill, context={"request": request})
        return Response(ser.data, status=status.HTTP_200_OK)

    @action(detail=True, methods=["post"], url_path="delete-template")
    def delete_template(self, request, pk=None):
        """删除团队模板文件。"""
        skill = self.get_object()
        if skill.template_file:
            try:
                skill.template_file.delete(save=False)
            except Exception:
                pass
        skill.template_file = None
        skill.template_columns = []
        skill.save()
        return Response({"detail": "模板已删除"}, status=status.HTTP_200_OK)

    @action(detail=False, methods=["post"], url_path="run")
    def run_skill(self, request):
        """通用 Skill 执行：用指定 Skill 的系统提示词处理输入文本，返回生成结果。

        供各业务模块（接口测试 / UI 自动化 / APP 自动化 / 性能测试 / AI 智能模式）
        直接调用，无需经过 Hermes Agent。逻辑与 Hermes 的 use_skill 工具一致。
        """
        skill_id = request.data.get("skill_id")
        input_text = request.data.get("input_text", "")
        if not skill_id:
            return Response({"detail": "skill_id 不能为空"}, status=status.HTTP_400_BAD_REQUEST)
        if not str(input_text).strip():
            return Response({"detail": "input_text 不能为空"}, status=status.HTTP_400_BAD_REQUEST)
        try:
            skill = TestCaseSkill.objects.get(pk=skill_id, is_active=True)
        except TestCaseSkill.DoesNotExist:
            return Response({"detail": f"Skill {skill_id} 不存在或未启用"}, status=status.HTTP_404_NOT_FOUND)

        model_config = skill.resolve_model_config()
        if not model_config:
            return Response(
                {"detail": "未配置可用的 AI 模型，请先在配置中心添加 writer 角色的模型"},
                status=status.HTTP_400_BAD_REQUEST,
            )
        system_prompt = skill.resolve_prompt_content()
        if not system_prompt.strip():
            return Response({"detail": "该 Skill 没有配置系统提示词"}, status=status.HTTP_400_BAD_REQUEST)

        try:
            logger.info("Skill 执行开始: skill=%s (%s), model=%s", skill.name, skill.skill_type, model_config.name)
            messages = [
                {"role": "system", "content": system_prompt},
                {"role": "user", "content": str(input_text)},
            ]
            response = async_to_sync(AIModelService.call_openai_compatible_api)(model_config, messages)
            content = response.get("choices", [{}])[0].get("message", {}).get("content", "")
            logger.info("Skill 执行完成: skill=%s, output_chars=%d", skill.name, len(content))
            return Response({
                "skill_id": skill.id,
                "skill_name": skill.name,
                "skill_type": skill.skill_type,
                "output_format": skill.output_format,
                "output": content,
            })
        except Exception as exc:
            logger.exception("Skill 执行失败: skill=%s", skill.name)
            return Response({"detail": f"Skill 执行失败: {exc}"}, status=status.HTTP_500_INTERNAL_ERROR)

    @action(detail=False, methods=["post"], url_path="save")
    def save_skill_case(self, request):
        """将 Skill 生成的用例文本保存为对应模块的真实用例/脚本对象。

        参数：module(必填)、content(必填)、name、project_id(可选，按模块语义为模块项目 id)、skill_id(可选)。
        映射规则见 skill_case_saver.save_case。
        """
        module = request.data.get("module")
        content = request.data.get("content", "")
        name = request.data.get("name", "")
        project_id = request.data.get("project_id") or None
        skill_id = request.data.get("skill_id")
        if not module:
            return Response({"detail": "module 不能为空"}, status=status.HTTP_400_BAD_REQUEST)
        if not str(content).strip():
            return Response({"detail": "content 不能为空"}, status=status.HTTP_400_BAD_REQUEST)
        skill = None
        if skill_id:
            skill = TestCaseSkill.objects.filter(pk=skill_id, is_active=True).first()
        try:
            from .skill_case_saver import save_case
            result = save_case(module, content, name, project_id, request.user, skill=skill)
            return Response(result, status=status.HTTP_201_CREATED)
        except ValueError as exc:
            return Response({"detail": str(exc)}, status=status.HTTP_400_BAD_REQUEST)
        except Exception as exc:
            logger.exception("保存 Skill 用例失败")
            return Response({"detail": f"保存失败: {exc}"}, status=status.HTTP_500_INTERNAL_ERROR)

    @action(detail=True, methods=["get"], url_path="export-package")
    def export_package(self, request, pk=None):
        """导出技能包：返回 .skill.zip（manifest.json + README.md + prompt.md + constraint.md + input/ + assets/ + references/ + scripts/ + resources/）。"""
        skill = self.get_object()
        try:
            from .skill_package import build_skill_package

            data = build_skill_package(skill)
        except Exception as exc:
            logger.exception("导出技能包失败")
            return Response({"detail": f"导出失败：{exc}"}, status=status.HTTP_500_INTERNAL_SERVER_ERROR)

        from django.utils.encoding import escape_uri_path
        safe_name = (skill.name or "skill").replace(" ", "_")
        filename = f"{safe_name}.skill.zip"
        resp = HttpResponse(data, content_type="application/zip")
        resp["Content-Disposition"] = f"attachment; filename*=UTF-8''{escape_uri_path(filename)}"
        return resp

    @action(detail=False, methods=["post"], url_path="import-package")
    def import_package(self, request):
        """导入技能包：上传 .skill.zip 或 .json，创建或更新 Skill。

        若已存在同名 Skill 则更新其内容（re-import 语义），同时替换包内文件。
        """
        uploaded = request.data.get("file")
        if not uploaded:
            return Response({"detail": "请上传 .skill.zip 或 .json 技能包文件"}, status=status.HTTP_400_BAD_REQUEST)
        try:
            from .skill_package import parse_skill_package, manifest_to_model_kwargs, save_artifact_files
            from .models import TestCaseSkill

            manifest, resources, artifact_files = parse_skill_package(uploaded)
            if not (manifest.get("name") or "").strip():
                return Response({"detail": "技能包缺少 name 字段"}, status=status.HTTP_400_BAD_REQUEST)

            kwargs, template_file_obj = manifest_to_model_kwargs(manifest, resources, TestCaseSkill)
            # created_by：导入者
            kwargs["created_by"] = request.user

            existing = TestCaseSkill.objects.filter(name=manifest["name"]).first()
            if existing:
                for k, v in kwargs.items():
                    if k == "created_by":
                        continue
                    setattr(existing, k, v)
                if template_file_obj is not None:
                    existing.template_file.save(template_file_obj.name, template_file_obj, save=False)
                existing.save()
                skill = existing
                created = False
            else:
                skill = TestCaseSkill(**kwargs)
                skill.save()
                if template_file_obj is not None:
                    skill.template_file.save(template_file_obj.name, template_file_obj, save=True)
                created = True

            # 保存包内文件
            save_artifact_files(skill, artifact_files)

            return Response(
                {
                    "detail": "导入成功" if created else "已更新同名技能",
                    "id": skill.id,
                    "name": skill.name,
                    "created": created,
                    "file_count": skill.artifacts.count(),
                },
                status=status.HTTP_201_CREATED if created else status.HTTP_200_OK,
            )
        except Exception as exc:
            logger.exception("导入技能包失败")
            return Response({"detail": f"导入失败：{exc}"}, status=status.HTTP_400_BAD_REQUEST)

    @action(detail=True, methods=["post"], url_path="add-artifact")
    def add_artifact(self, request, pk=None):
        """向 Skill 包中添加/更新一个文件。"""
        skill = self.get_object()
        path = request.data.get("path") or request.data.get("filename")
        text_content = request.data.get("text_content", "")
        artifact_type = request.data.get("artifact_type", "other")
        uploaded = request.FILES.get("file")
        if not path:
            return Response({"detail": "path 不能为空"}, status=status.HTTP_400_BAD_REQUEST)
        from .models import SkillArtifact
        from django.core.files.base import ContentFile

        artifact, _ = SkillArtifact.objects.update_or_create(
            skill=skill, path=path,
            defaults={"artifact_type": artifact_type, "text_content": ""},
        )
        if uploaded:
            artifact.is_binary = not uploaded.content_type or not uploaded.content_type.startswith("text/")
            artifact.file.save(uploaded.name, uploaded, save=False)
            artifact.text_content = ""
        else:
            artifact.is_binary = False
            artifact.text_content = text_content
        artifact.save()
        from .serializers import SkillArtifactSerializer
        return Response(SkillArtifactSerializer(artifact, context={"request": request}).data)

    @action(detail=True, methods=["post"], url_path="remove-artifact")
    def remove_artifact(self, request, pk=None):
        """删除 Skill 包中的一个文件。"""
        skill = self.get_object()
        artifact_id = request.data.get("artifact_id")
        path = request.data.get("path")
        from .models import SkillArtifact
        qs = skill.artifacts.all()
        if artifact_id:
            qs = qs.filter(pk=artifact_id)
        elif path:
            qs = qs.filter(path=path)
        else:
            return Response({"detail": "artifact_id 或 path 必填"}, status=status.HTTP_400_BAD_REQUEST)
        count, _ = qs.delete()
        return Response({"detail": f"已删除 {count} 个文件"})


def _parse_template_columns(uploaded_file) -> list:
    """从上传的 Excel/CSV 文件解析首行列名。

    支持 .xlsx（openpyxl）/ .xls（xlrd 可选）/ .csv（utf-8/sig/gbk）。
    返回去重后的列名列表。
    """
    import csv as _csv
    import io
    import os

    name = (getattr(uploaded_file, "name", "") or "").lower()
    content = uploaded_file.read()
    columns: list = []

    if name.endswith(".csv"):
        # 探测编码
        for enc in ("utf-8-sig", "utf-8", "gbk"):
            try:
                text = content.decode(enc)
                break
            except Exception:
                text = None
        if text is None:
            text = content.decode("utf-8", errors="ignore")
        reader = _csv.reader(io.StringIO(text))
        for row in reader:
            if row:
                columns = [c.strip() for c in row if c and c.strip()]
                break
    else:
        # Excel：openpyxl
        from openpyxl import load_workbook

        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        ws = wb.active
        for row in ws.iter_rows(min_row=1, max_row=1, values_only=True):
            if row:
                columns = [str(c).strip() for c in row if c is not None and str(c).strip()]
            break
        wb.close()

    # 去重，保留顺序
    seen = set()
    result = []
    for c in columns:
        if c not in seen:
            seen.add(c)
            result.append(c)
    return result


def _swagger_to_requirement_text(swagger: dict) -> str:
    """把 Swagger/OpenAPI 字典转成可读的接口需求文本（供用例生成使用）。"""
    lines = []
    info = swagger.get("info") or {}
    if info.get("title"):
        lines.append(f"接口文档：{info.get('title')}")
    if info.get("version"):
        lines.append(f"版本：{info.get('version')}")
    paths = swagger.get("paths") or {}
    if not paths:
        lines.append("（未找到 paths 定义）")
        return "\n".join(lines)

    method_cn = {
        "get": "查询", "post": "新增", "put": "更新(全量)",
        "patch": "更新(部分)", "delete": "删除", "head": "HEAD", "options": "OPTIONS",
    }
    for path, methods in paths.items():
        if not isinstance(methods, dict):
            continue
        for method, op in methods.items():
            if method.lower() not in ("get", "post", "put", "patch", "delete", "head", "options"):
                continue
            op = op or {}
            summary = op.get("summary") or op.get("description") or ""
            line = f"- [{method.upper()}] {path}  {method_cn.get(method.lower(), method)}"
            if summary:
                line += f"：{summary}"
            lines.append(line)
            # 参数概要
            params = op.get("parameters") or []
            for p in params:
                p = p or {}
                pname = p.get("name")
                pdesc = p.get("description") or ""
                pin = p.get("in", "")
                if pname:
                    lines.append(f"    - 参数 {pname}（{pin}）：{pdesc}".rstrip())
            # 请求体
            req_body = (op.get("requestBody") or {}).get("content") or {}
            if req_body:
                lines.append("    - 请求体：详见接口 schema")
    return "\n".join(lines)


class ConfigStatusViewSet(viewsets.ViewSet):
    permission_classes = [IsAuthenticated]

    def list(self, request):
        data: Dict[str, Any] = {
            "ai_model_config_active": AIModelConfig.objects.filter(is_active=True).exists(),
            "prompt_config_active": PromptConfig.objects.filter(is_active=True).exists(),
            "generation_config_active": GenerationConfig.objects.filter(is_active=True).exists(),
        }
        return Response(data)

    @action(detail=False, methods=["get"], url_path="check")
    def check(self, request):
        """
        兼容旧前端页面：返回细粒度配置状态（/requirement-analysis/config/check/）。

        返回结构与旧版页面 `RequirementAnalysisView.vue` 对齐：
        - writer_model / reviewer_model：是否已配置、是否启用
        - writer_prompt / reviewer_prompt：是否已配置、是否启用
        - generation_config：是否已配置 + 默认输出模式 + 是否启用自动评审
        """

        def _model_status(role: str) -> Dict[str, Any]:
            qs = AIModelConfig.objects.filter(role=role)
            configured = qs.exists()
            enabled = qs.filter(is_active=True).exists()
            active = qs.filter(is_active=True).order_by("-updated_at").first()
            return {
                "configured": configured,
                "enabled": enabled,
                "active_id": getattr(active, "id", None),
                "active_name": getattr(active, "name", None),
            }

        def _prompt_status(prompt_type: str) -> Dict[str, Any]:
            qs = PromptConfig.objects.filter(prompt_type=prompt_type)
            configured = qs.exists()
            enabled = qs.filter(is_active=True).exists()
            active = qs.filter(is_active=True).order_by("-updated_at").first()
            return {
                "configured": configured,
                "enabled": enabled,
                "active_id": getattr(active, "id", None),
                "active_name": getattr(active, "name", None),
            }

        gc_qs = GenerationConfig.objects.all()
        gc_configured = gc_qs.exists()
        gc_active = gc_qs.filter(is_active=True).order_by("-updated_at").first() or gc_qs.order_by("-updated_at").first()

        data: Dict[str, Any] = {
            "writer_model": _model_status("writer"),
            "writer_prompt": _prompt_status("writer"),
            "reviewer_model": _model_status("reviewer"),
            "reviewer_prompt": _prompt_status("reviewer"),
            "generation_config": {
                "configured": gc_configured,
                "active_id": getattr(gc_active, "id", None),
                "default_output_mode": getattr(gc_active, "default_output_mode", None),
                "enable_auto_review": getattr(gc_active, "enable_auto_review", None),
            },
        }
        return Response(data)


@api_view(["POST"])
@permission_classes([IsAuthenticated])
def analyze_text(request):
    """
    对纯文本进行需求分析（不落库/不生成 BusinessRequirement 记录）。
    """
    text = (request.data or {}).get("text") or ""
    title = (request.data or {}).get("title") or ""
    if not text.strip():
        return Response({"detail": "text 不能为空"}, status=status.HTTP_400_BAD_REQUEST)

    try:
        result = async_to_sync(AIService.analyze_requirements)(text, title)
        return Response(result, status=status.HTTP_200_OK)
    except Exception as e:
        logger.exception("文本分析失败")
        return Response({"detail": str(e)}, status=status.HTTP_500_INTERNAL_SERVER_ERROR)


@api_view(["POST"])
@permission_classes([IsAuthenticated])
def upload_and_analyze(request):
    """
    上传文档并立即执行分析（同步等待完成，适合开发/小文档）。
    """
    ser = DocumentUploadSerializer(data=request.data, context={"request": request})
    ser.is_valid(raise_exception=True)
    doc = ser.save()

    try:
        analysis = async_to_sync(RequirementAnalysisService.process_document_analysis)(doc)
        return Response(
            {
                "document": RequirementDocumentSerializer(doc).data,
                "analysis": RequirementAnalysisSerializer(analysis).data,
            },
            status=status.HTTP_201_CREATED,
        )
    except Exception as e:
        logger.exception("上传并分析失败")
        return Response({"detail": str(e)}, status=status.HTTP_500_INTERNAL_SERVER_ERROR)
